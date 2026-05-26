from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.application.interfaces.report_exporter_interface import ReportExporterInterface
from app.domain.entities.change_report import ChangeReport
from app.domain.value_objects.asset_type import AssetType
from app.domain.value_objects.severity_level import SeverityLevel
from app.shared.filename_utils import generate_unique_report_path
from app.shared.logger import logger


class ExcelReportExporter(ReportExporterInterface):
    def export(self, report: ChangeReport, output_dir: Path) -> Path:
        """
        Exporta el reporte consolidado a un archivo Excel (.xlsx) estructurado en 9 hojas
        con diseño y colores premium.
        """
        output_path = generate_unique_report_path(
            output_dir, "attack_surface_changes_report", ".xlsx"
        )

        # 1. Preparar datos para cada pestaña

        # Hoja 1: Resumen
        resumen_df = pd.DataFrame(
            [
                {
                    "total_cambios": report.total_changes,
                    "cambios_criticos": report.critical_changes,
                    "cambios_altos": report.high_changes,
                    "cambios_medios": report.medium_changes,
                    "cambios_bajos": report.low_changes,
                    "cambios_informativos": report.info_changes,
                    "riesgo_general": f"{report.riesgo_general:.2f} ({report.risk_level_general.value})",
                    "fecha_generacion": report.generated_at,
                }
            ]
        )

        # Hoja 2: Cambios Detectados
        cambios_rows = []
        for c in report.changes:
            cambios_rows.append(
                {
                    "change_id": c.change_id,
                    "asset_type": c.asset_type.value,
                    "change_type": c.change_type.value,
                    "asset_identifier": c.asset_identifier,
                    "previous_value": c.previous_value,
                    "current_value": c.current_value,
                    "severity": c.severity.value,
                    "risk_score": c.risk_score,
                    "risk_level": c.risk_level.value,
                    "description": c.description,
                    "impact": c.impact,
                    "recommendation": c.recommendation,
                    "detected_at": c.detected_at,
                }
            )
        cambios_df = (
            pd.DataFrame(cambios_rows)
            if cambios_rows
            else pd.DataFrame(
                columns=[
                    "change_id",
                    "asset_type",
                    "change_type",
                    "asset_identifier",
                    "previous_value",
                    "current_value",
                    "severity",
                    "risk_score",
                    "risk_level",
                    "description",
                    "impact",
                    "recommendation",
                    "detected_at",
                ]
            )
        )

        # Hoja 3: Puertos
        ports_rows = []
        # Hoja 4: Subdominios
        sub_rows = []
        # Hoja 5: Servicios
        svc_rows = []
        # Hoja 6: Headers
        header_rows = []
        # Hoja 7: SSL_TLS
        ssl_rows = []
        # Hoja 8: Recomendaciones
        reco_rows = []

        # Recorrer los cambios crudos vinculados al reporte en el caso de uso
        raw_changes = getattr(report, "raw_changes", [])

        for idx, rc in enumerate(raw_changes):
            details = rc.get("details", {})
            entity = rc.get("entity")

            severity_str = entity.severity.value if entity else "INFO"
            recommendation_str = entity.recommendation if entity else ""

            # Pestaña 3: Puertos
            if rc["asset_type"] == AssetType.PORT:
                ports_rows.append(
                    {
                        "target": details.get("target"),
                        "ip": details.get("ip"),
                        "port": details.get("port"),
                        "previous_status": details.get("previous_status"),
                        "current_status": details.get("current_status"),
                        "previous_service": details.get("previous_service"),
                        "current_service": details.get("current_service"),
                        "severity": severity_str,
                        "recommendation": recommendation_str,
                    }
                )

            # Pestaña 4: Subdominios
            elif rc["asset_type"] == AssetType.SUBDOMAIN:
                sub_rows.append(
                    {
                        "domain": details.get("domain"),
                        "subdomain": details.get("subdomain"),
                        "previous_ip": details.get("previous_ip"),
                        "current_ip": details.get("current_ip"),
                        "previous_status": details.get("previous_status"),
                        "current_status": details.get("current_status"),
                        "severity": severity_str,
                        "recommendation": recommendation_str,
                    }
                )

            # Pestaña 5: Servicios
            elif rc["asset_type"] == AssetType.SERVICE:
                svc_rows.append(
                    {
                        "host": details.get("host"),
                        "port": details.get("port"),
                        "previous_service": details.get("previous_service"),
                        "current_service": details.get("current_service"),
                        "previous_version": details.get("previous_version"),
                        "current_version": details.get("current_version"),
                        "severity": severity_str,
                        "recommendation": recommendation_str,
                    }
                )

            # Pestaña 6: Headers
            elif rc["asset_type"] == AssetType.HEADER:
                header_rows.append(
                    {
                        "url": details.get("url"),
                        "header": details.get("header"),
                        "previous_value": details.get("previous_value"),
                        "current_value": details.get("current_value"),
                        "change_type": rc["change_type"].value,
                        "severity": severity_str,
                        "recommendation": recommendation_str,
                    }
                )

            # Pestaña 7: SSL_TLS
            elif rc["asset_type"] == AssetType.SSL_TLS:
                ssl_rows.append(
                    {
                        "domain": details.get("domain"),
                        "check_name": details.get("check_name"),
                        "previous_value": details.get("previous_value"),
                        "current_value": details.get("current_value"),
                        "change_type": rc["change_type"].value,
                        "severity": severity_str,
                        "recommendation": recommendation_str,
                    }
                )

            # Agregar a recomendaciones prioritarias si la severidad es MEDIA o superior
            if entity and entity.severity in [
                SeverityLevel.CRITICAL,
                SeverityLevel.HIGH,
                SeverityLevel.MEDIUM,
            ]:
                priority_map = {
                    SeverityLevel.CRITICAL: 1,
                    SeverityLevel.HIGH: 2,
                    SeverityLevel.MEDIUM: 3,
                }
                reco_rows.append(
                    {
                        "prioridad": priority_map.get(entity.severity, 4),
                        "asset_type": entity.asset_type.value,
                        "asset_identifier": entity.asset_identifier,
                        "problema": entity.description,
                        "recomendacion": entity.recommendation,
                    }
                )

        # Dataframes correspondientes
        ports_df = (
            pd.DataFrame(ports_rows)
            if ports_rows
            else pd.DataFrame(
                columns=[
                    "target",
                    "ip",
                    "port",
                    "previous_status",
                    "current_status",
                    "previous_service",
                    "current_service",
                    "severity",
                    "recommendation",
                ]
            )
        )
        sub_df = (
            pd.DataFrame(sub_rows)
            if sub_rows
            else pd.DataFrame(
                columns=[
                    "domain",
                    "subdomain",
                    "previous_ip",
                    "current_ip",
                    "previous_status",
                    "current_status",
                    "severity",
                    "recommendation",
                ]
            )
        )
        svc_df = (
            pd.DataFrame(svc_rows)
            if svc_rows
            else pd.DataFrame(
                columns=[
                    "host",
                    "port",
                    "previous_service",
                    "current_service",
                    "previous_version",
                    "current_version",
                    "severity",
                    "recommendation",
                ]
            )
        )
        header_df = (
            pd.DataFrame(header_rows)
            if header_rows
            else pd.DataFrame(
                columns=[
                    "url",
                    "header",
                    "previous_value",
                    "current_value",
                    "change_type",
                    "severity",
                    "recommendation",
                ]
            )
        )
        ssl_df = (
            pd.DataFrame(ssl_rows)
            if ssl_rows
            else pd.DataFrame(
                columns=[
                    "domain",
                    "check_name",
                    "previous_value",
                    "current_value",
                    "change_type",
                    "severity",
                    "recommendation",
                ]
            )
        )

        # Ordenar recomendaciones por prioridad
        reco_df = (
            pd.DataFrame(reco_rows)
            if reco_rows
            else pd.DataFrame(
                columns=["prioridad", "asset_type", "asset_identifier", "problema", "recomendacion"]
            )
        )
        if not reco_df.empty:
            reco_df = reco_df.sort_values(by="prioridad")

        # Hoja 9: Errores
        errors_df = (
            pd.DataFrame(getattr(report, "errors", []))
            if getattr(report, "errors", [])
            else pd.DataFrame(columns=["archivo", "tipo_error", "mensaje_error", "fecha"])
        )

        # 2. Escribir a Excel usando pandas ExcelWriter
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            resumen_df.to_excel(writer, sheet_name="Resumen", index=False)
            cambios_df.to_excel(writer, sheet_name="Cambios Detectados", index=False)
            ports_df.to_excel(writer, sheet_name="Puertos", index=False)
            sub_df.to_excel(writer, sheet_name="Subdominios", index=False)
            svc_df.to_excel(writer, sheet_name="Servicios", index=False)
            header_df.to_excel(writer, sheet_name="Headers", index=False)
            ssl_df.to_excel(writer, sheet_name="SSL_TLS", index=False)
            reco_df.to_excel(writer, sheet_name="Recomendaciones", index=False)
            errors_df.to_excel(writer, sheet_name="Errores", index=False)

            # Estilizar las hojas
            workbook = writer.book

            # Estilos corporativos
            header_fill = PatternFill(
                start_color="1B365D", end_color="1B365D", fill_type="solid"
            )  # Azul oscuro ciberseguridad
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_regular = Font(name="Segoe UI", size=10)

            # Bordes sutiles
            thin_border = Border(
                left=Side(style="thin", color="E0E0E0"),
                right=Side(style="thin", color="E0E0E0"),
                top=Side(style="thin", color="E0E0E0"),
                bottom=Side(style="thin", color="E0E0E0"),
            )

            # Fills para severidades
            severity_fills = {
                "CRITICAL": PatternFill(
                    start_color="FFD6D6", end_color="FFD6D6", fill_type="solid"
                ),  # Rojo claro
                "HIGH": PatternFill(
                    start_color="FFEAD2", end_color="FFEAD2", fill_type="solid"
                ),  # Naranja claro
                "MEDIUM": PatternFill(
                    start_color="FFFDD0", end_color="FFFDD0", fill_type="solid"
                ),  # Amarillo claro
                "LOW": PatternFill(
                    start_color="EAF2F8", end_color="EAF2F8", fill_type="solid"
                ),  # Azul claro
                "INFO": PatternFill(
                    start_color="EAFAF1", end_color="EAFAF1", fill_type="solid"
                ),  # Verde claro
            }

            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                worksheet.views.sheetView[0].showGridLines = True

                # Fila de cabecera
                for col_idx in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                    cell.border = thin_border

                # Configurar alto de cabecera
                worksheet.row_dimensions[1].height = 28

                # Celdas de datos
                for row_idx in range(2, worksheet.max_row + 1):
                    worksheet.row_dimensions[row_idx].height = 20

                    # Identificar la columna de severidad para pintar la fila
                    severity_val = None
                    for col_idx in range(1, worksheet.max_column + 1):
                        cell = worksheet.cell(row=row_idx, column=col_idx)
                        cell.font = font_regular
                        cell.border = thin_border

                        # Buscar si esta columna representa severidad
                        header_val = str(worksheet.cell(row=1, column=col_idx).value).lower()
                        if "severity" in header_val:
                            severity_val = str(cell.value).upper()

                    # Si tiene una severidad identificada, aplicar patrón de color a toda la fila
                    if severity_val in severity_fills:
                        for col_idx in range(1, worksheet.max_column + 1):
                            cell = worksheet.cell(row=row_idx, column=col_idx)
                            cell.fill = severity_fills[severity_val]

                # Ajustar anchos de columnas automáticamente
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    # Ancho proporcional con límites razonables
                    worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        logger.info(f"Reporte Excel generado correctamente en: {output_path.name}")
        return output_path
